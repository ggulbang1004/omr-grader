"""채점 + 엑셀 리포트 (문항 수 가변 대응)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="DCE6F1")
HDR_FONT = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ERR_FILL = PatternFill("solid", fgColor="FCE4EC")


def score_answers(read_map, key, points):
    """read_map: {q: {'mark': [...]}}. -> (total, per_q)."""
    total = 0
    per_q = {}
    for q in sorted(read_map):
        if q == "_thresh":
            continue
        m = read_map[q]["mark"]
        if q not in key:
            per_q[q] = dict(mark="-", pts=0)
            continue
        if len(m) == 0:
            per_q[q] = dict(mark="미기재", pts=0)
        elif len(m) > 1:
            per_q[q] = dict(mark="복수표기", pts=0)
        elif m[0] == key[q]:
            per_q[q] = dict(mark="O", pts=points.get(q, 2))
            total += points.get(q, 2)
        else:
            per_q[q] = dict(mark="X", pts=0)
    return total, per_q


def build_workbook(students, key, points, period=""):
    """students: list of dict(exam_no, name, total, per_q, errors).

    errors: list of (구분, 내용).
    """
    questions = sorted(key.keys())
    wb = Workbook()

    ws = wb.active
    ws.title = "개인별점수"
    headers = ["수험번호", "성명", "점수"] + [str(q) for q in questions]
    ws.append(headers)
    for s in students:
        row = [s.get("exam_no", ""), s.get("name", ""), s.get("total", 0)]
        row += [s["per_q"].get(q, {}).get("mark", "-") for q in questions]
        ws.append(row)
    _style_header(ws, len(headers))
    # 오류 행 하이라이트
    err_rows = set()
    for i, s in enumerate(students, start=2):
        if s.get("errors"):
            err_rows.add(i)
    for r in err_rows:
        for c in range(1, 4):
            ws.cell(row=r, column=c).fill = ERR_FILL

    ws2 = wb.create_sheet("문항분석")
    ws2.append(["문항", "정답", "배점", "응시자수", "정답자수", "정답률(%)"])
    n = len(students)
    for q in questions:
        correct = sum(1 for s in students
                      if s["per_q"].get(q, {}).get("mark") == "O")
        rate = round(correct / n * 100, 1) if n else 0
        ws2.append([q, key[q], points.get(q, 2), n, correct, rate])
    _style_header(ws2, 6)

    ws3 = wb.create_sheet("오류리스트")
    ws3.append(["수험번호", "성명", "구분", "내용"])
    for s in students:
        for e in s.get("errors", []):
            ws3.append([s.get("exam_no", ""), s.get("name", ""),
                        e[0] if len(e) > 0 else "", e[1] if len(e) > 1 else ""])
    _style_header(ws3, 4)
    if ws3.max_row == 1:
        ws3.append(["-", "-", "없음", "오류 없음"])

    ws4 = wb.create_sheet("정답및배점")
    ws4.append(["문항", "정답", "배점"])
    for q in questions:
        ws4.append([q, key[q], points.get(q, 2)])
    _style_header(ws4, 3)

    for w in wb.worksheets:
        w.freeze_panes = "A2"
        w.sheet_properties.pageSetUpPr.fitToPage = True
    _autowidth(ws, max_w=12)
    if period:
        wb.properties.title = f"가채점 결과 - {period}"
    return wb


def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def _autowidth(ws, max_w=14):
    for col in ws.columns:
        mx = 0
        for cell in col:
            v = cell.value
            mx = max(mx, len(str(v)) if v is not None else 0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(
            max(mx + 2, 8), max_w)
