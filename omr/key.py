"""정답·배점 엑셀 로더. 컬럼명 변형 허용."""
import re
from openpyxl import load_workbook

HEADER_MAP = {
    "q": ["문항", "번호", "문항번호", "q", "no"],
    "answer": ["정답", "답", "정답번호", "answer"],
    "points": ["배점", "점수", "points", "score"],
}


def _norm(h):
    return re.sub(r"\s+", "", str(h or "")).lower()


def load_key(xlsx_path, sheet=None):
    """-> (key {q: 정답}, points {q: 배점}). 정답은 1-5 숫자."""
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    rows = list(ws.iter_rows(values_only=True))
    header_idx, cols = None, {}
    for i, r in enumerate(rows[:5]):
        c = {}
        for idx, h in enumerate(r):
            nh = _norm(h)
            for k, vs in HEADER_MAP.items():
                if nh in [_norm(v) for v in vs] and k not in c:
                    c[k] = idx
        if "q" in c and "answer" in c:
            header_idx, cols = i, c
            break
    if header_idx is None:
        raise ValueError("정답 엑셀에서 문항/정답 컬럼을 찾지 못했습니다.")
    key, points = {}, {}
    for r in rows[header_idx + 1:]:
        if r[cols["q"]] is None or r[cols["answer"]] is None:
            continue
        try:
            q = int(r[cols["q"]])
        except (ValueError, TypeError):
            continue
        a = str(r[cols["answer"]]).strip()
        m = re.search(r"[1-5]", a)
        if not m:
            continue
        key[q] = int(m.group(0))
        p = r[cols["points"]] if "points" in cols else 2
        try:
            points[q] = float(p)
        except (ValueError, TypeError):
            points[q] = 2
    if not key:
        raise ValueError("정답 데이터가 비어 있습니다.")
    return key, points
